import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

def build_presentation():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Professional Color Palette
    PRIMARY = RGBColor(15, 23, 42)        # Deep Slate / Navy
    BRAND = RGBColor(37, 99, 235)         # Royal Blue
    ACCENT = RGBColor(14, 165, 233)       # Sky Cyan
    BG_LIGHT = RGBColor(248, 250, 252)    # Slate 50
    CARD_BG = RGBColor(255, 255, 255)     # White
    CARD_BORDER = RGBColor(203, 213, 225) # Slate 300
    TEXT_DARK = RGBColor(15, 23, 42)      # Slate 900
    TEXT_MUTED = RGBColor(71, 85, 105)    # Slate 600
    DOT_RED = RGBColor(239, 68, 68)
    DOT_YELLOW = RGBColor(245, 158, 11)
    DOT_GREEN = RGBColor(16, 185, 129)

    img_dir = r"c:\Users\A D M I N\Desktop\project\mini project\ppt_images"

    def set_bg(slide, color=BG_LIGHT):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
        bg.fill.solid()
        bg.fill.fore_color.rgb = color
        bg.line.color.rgb = color
        return bg

    def add_browser_frame(slide, left, top, width, height, image_path, url_text="http://localhost/crms/"):
        # Browser outer frame
        frame = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        frame.fill.solid()
        frame.fill.fore_color.rgb = CARD_BG
        frame.line.color.rgb = CARD_BORDER
        frame.line.width = Pt(1.5)

        # Browser toolbar bar
        toolbar_h = Inches(0.4)
        toolbar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, toolbar_h)
        toolbar.fill.solid()
        toolbar.fill.fore_color.rgb = RGBColor(241, 245, 249)
        toolbar.line.color.rgb = CARD_BORDER
        toolbar.line.width = Pt(1)

        # 3 macOS-style window dots
        dots = [DOT_RED, DOT_YELLOW, DOT_GREEN]
        for idx, col in enumerate(dots):
            dot = slide.shapes.add_shape(MSO_SHAPE.OVAL, left + Inches(0.15 + idx * 0.18), top + Inches(0.12), Inches(0.14), Inches(0.14))
            dot.fill.solid()
            dot.fill.fore_color.rgb = col
            dot.line.fill.background()

        # URL address pill
        url_pill = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left + Inches(0.85), top + Inches(0.08), width - Inches(1.1), Inches(0.24))
        url_pill.fill.solid()
        url_pill.fill.fore_color.rgb = CARD_BG
        url_pill.line.color.rgb = RGBColor(226, 232, 240)
        
        p = url_pill.text_frame.paragraphs[0]
        p.text = "🔒 " + url_text
        p.font.size = Pt(9)
        p.font.color.rgb = TEXT_MUTED
        p.font.name = 'Consolas'

        # Insert Screenshot Image
        img_top = top + toolbar_h + Inches(0.02)
        img_h = height - toolbar_h - Inches(0.04)
        img_w = width - Inches(0.04)
        slide.shapes.add_picture(image_path, left + Inches(0.02), img_top, img_w, img_h)

    # -------------------------------------------------------------
    # SLIDE 1: Title & Hero Showcase (With Dashboard Screenshot)
    # -------------------------------------------------------------
    s1 = prs.slides.add_slide(blank_layout)
    set_bg(s1, PRIMARY)

    # Left content
    tbox = s1.shapes.add_textbox(Inches(0.8), Inches(1.4), Inches(5.2), Inches(5.0))
    tf1 = tbox.text_frame
    tf1.word_wrap = True

    p = tf1.paragraphs[0]
    p.text = "CAMPUS RECRUITMENT SYSTEM"
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = ACCENT
    p.font.name = 'Arial'

    p2 = tf1.add_paragraph()
    p2.text = "NexusCRMS"
    p2.font.size = Pt(44)
    p2.font.bold = True
    p2.font.color.rgb = RGBColor(255, 255, 255)
    p2.font.name = 'Arial'
    p2.space_before = Pt(8)

    p3 = tf1.add_paragraph()
    p3.text = "Project Demonstration & Live Website Walkthrough"
    p3.font.size = Pt(20)
    p3.font.color.rgb = RGBColor(191, 219, 254)
    p3.font.name = 'Arial'
    p3.space_before = Pt(8)

    p4 = tf1.add_paragraph()
    p4.text = "An interactive, automated platform connecting college placement cells, recruiting companies, and eligible students."
    p4.font.size = Pt(14)
    p4.font.color.rgb = RGBColor(148, 163, 184)
    p4.space_before = Pt(18)

    p5 = tf1.add_paragraph()
    p5.text = "• Role-Based Access (Admin & Student)\n• Real-Time CGPA & Branch Eligibility Engine\n• 4-Stage Visual Application Pipeline\n• Built on PHP (MVC) + MySQL + Bootstrap"
    p5.font.size = Pt(13)
    p5.font.color.rgb = RGBColor(203, 213, 225)
    p5.space_before = Pt(14)

    # Right side: Hero browser frame with Dashboard screenshot
    dash_img = os.path.join(img_dir, "01_admin_dashboard.png")
    add_browser_frame(s1, Inches(6.2), Inches(1.2), Inches(6.4), Inches(5.2), dash_img, "localhost/crms/admin/dashboard")

    # -------------------------------------------------------------
    # SHOWCASE SLIDES (1 Slide per Screen with Screenshot & Features)
    # -------------------------------------------------------------
    screens = [
        {
            "cat": "ADMINISTRATOR MODULE",
            "title": "Admin Placement Dashboard & Live KPIs",
            "url": "localhost/crms/admin/dashboard",
            "img": "01_admin_dashboard.png",
            "points": [
                ("Real-Time KPI Cards", "Live counters for Total Placed Candidates, Active Job Drives, Partner Companies, and Placement %."),
                ("Recruitment Activity Feed", "Monitors live application submissions and urgent stage-closing drives."),
                ("Fast Navigation Drawer", "Instant 1-click access to candidate pools, drive creation, company lists, and analytics.")
            ]
        },
        {
            "cat": "STUDENT MANAGEMENT",
            "title": "Student Directory & Verification Records",
            "url": "localhost/crms/admin/students",
            "img": "02_manage_students.png",
            "points": [
                ("Comprehensive Student Profiles", "Displays verified academic records: Roll number, Course/Branch, Semester, and CGPA."),
                ("Advanced Branch Filtering", "Quickly filter eligible pools across B.Tech CSE, IT, ECE, MCA, or Mechanical."),
                ("Resume & Contact Links", "Direct access to uploaded GitHub/PDF resumes and student contact channels.")
            ]
        },
        {
            "cat": "CORPORATE PARTNERS",
            "title": "Partner Companies & Recruiter Directory",
            "url": "localhost/crms/admin/companies",
            "img": "03_manage_companies.png",
            "points": [
                ("Corporate Directory (CRUD)", "Manage recruiting profiles for Google, Microsoft, AWS, TCS, Infosys, and startups."),
                ("Recruiter Contact Database", "Stores designated HR point of contact, corporate emails, phone numbers, and official websites."),
                ("Hiring History & Tiering", "Categorize partners into Tier-1 and mass hiring categories with active opening indicators.")
            ]
        },
        {
            "cat": "PLACEMENT DRIVES",
            "title": "Job Drives & Placement Drive Management",
            "url": "localhost/crms/admin/jobs",
            "img": "04_manage_jobs.png",
            "points": [
                ("Drive Configuration", "Create openings specifying Job Title, Job Type (Full-time / Internship), and CTC Salary Packages."),
                ("Strict Eligibility Criteria", "Set minimum CGPA cut-off, allowed degrees, eligible semesters, and required skill tags."),
                ("Deadlines & Status Control", "Enforces automated application deadlines and toggle active/closed drive status.")
            ]
        },
        {
            "cat": "APPLICATION PIPELINE",
            "title": "Applications Review & Candidate Roster",
            "url": "localhost/crms/admin/applications",
            "img": "05_applications.png",
            "points": [
                ("Centralized Applications Feed", "Real-time list of all submissions filtered by drive and hiring department."),
                ("Status Action Controls", "Easily update candidates to Under Review, Shortlisted, Selected, or Rejected."),
                ("Coordinator Notes", "Log confidential interview comments and technical round evaluation scores.")
            ]
        },
        {
            "cat": "SELECTION & HIRING",
            "title": "Shortlisting Board & Offer Management",
            "url": "localhost/crms/admin/shortlisting",
            "img": "06_shortlisting.png",
            "points": [
                ("Interactive Selection Board", "Drive-by-drive evaluation interface for placement coordinators."),
                ("Offer Letter Allocation", "Assign official institutional offer reference numbers (e.g. MSFT-2026-INT-0441)."),
                ("Remarks & Feedback", "Record special packaging terms, PPO agreements, and joining remarks.")
            ]
        },
        {
            "cat": "STUDENT PORTAL",
            "title": "Student Dashboard & Visual Application Tracker",
            "url": "localhost/crms/student/dashboard",
            "img": "07_student_dashboard.png",
            "points": [
                ("Personalized Student View", "Shows academic standing, current semester, verified CGPA, and skill badges."),
                ("4-Stage Visual Stepper", "Live color-coded status bar: Applied -> Under Review -> Shortlisted -> Selected!"),
                ("Instant Offer Recognition", "Direct view of received placement offers and coordinator recommendations.")
            ]
        },
        {
            "cat": "OPPORTUNITIES CATALOG",
            "title": "Smart Eligibility Engine & Job Catalog",
            "url": "localhost/crms/student/opportunities",
            "img": "08_student_catalog.png",
            "points": [
                ("Automated Eligibility Checks", "Real-time algorithmic check against student CGPA, branch, and semester."),
                ("Visual Status Indicators", "Green badge for Eligible drives (with skill match %); Red badge with exact disqualification reason."),
                ("1-Click Fast Apply", "Instant submission with guard checks preventing duplicate submissions.")
            ]
        }
    ]

    for item in screens:
        slide = prs.slides.add_slide(blank_layout)
        set_bg(slide, BG_LIGHT)

        # Category
        cat_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.7), Inches(0.35))
        tf_cat = cat_box.text_frame
        p_c = tf_cat.paragraphs[0]
        p_c.text = item["cat"]
        p_c.font.size = Pt(11)
        p_c.font.bold = True
        p_c.font.color.rgb = BRAND
        p_c.font.name = 'Arial'

        # Main Title
        t_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.7), Inches(11.7), Inches(0.6))
        tf_t = t_box.text_frame
        p_t = tf_t.paragraphs[0]
        p_t.text = item["title"]
        p_t.font.size = Pt(24)
        p_t.font.bold = True
        p_t.font.color.rgb = PRIMARY
        p_t.font.name = 'Arial'

        # Left Column: 3 Feature Cards
        col_w = Inches(3.8)
        for p_idx, (p_title, p_desc) in enumerate(item["points"]):
            c_top = Inches(1.5 + p_idx * 1.8)
            card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), c_top, col_w, Inches(1.65))
            card.fill.solid()
            card.fill.fore_color.rgb = CARD_BG
            card.line.color.rgb = CARD_BORDER
            card.line.width = Pt(1)

            tf_card = card.text_frame
            tf_card.word_wrap = True
            tf_card.margin_left = Inches(0.2)
            tf_card.margin_right = Inches(0.2)
            tf_card.margin_top = Inches(0.18)

            p_h = tf_card.paragraphs[0]
            p_h.text = f"{p_idx + 1}. {p_title}"
            p_h.font.size = Pt(14)
            p_h.font.bold = True
            p_h.font.color.rgb = BRAND

            p_b = tf_card.add_paragraph()
            p_b.text = p_desc
            p_b.font.size = Pt(11)
            p_b.font.color.rgb = TEXT_MUTED
            p_b.space_before = Pt(6)

        # Right Column: Big Website Browser Frame with Screenshot
        frame_left = Inches(4.85)
        frame_top = Inches(1.5)
        frame_w = Inches(7.7)
        frame_h = Inches(5.45)
        img_file = os.path.join(img_dir, item["img"])
        add_browser_frame(slide, frame_left, frame_top, frame_w, frame_h, img_file, item["url"])

    # -------------------------------------------------------------
    # SLIDE 10: Conclusion & Technology Stack (Dark Theme)
    # -------------------------------------------------------------
    s_end = prs.slides.add_slide(blank_layout)
    set_bg(s_end, PRIMARY)

    # Title
    t_end = s_end.shapes.add_textbox(Inches(0.8), Inches(1.0), Inches(11.7), Inches(1.2))
    tf_e = t_end.text_frame
    p_e = tf_e.paragraphs[0]
    p_e.text = "NexusCRMS: Technology & Architecture Summary"
    p_e.font.size = Pt(32)
    p_e.font.bold = True
    p_e.font.color.rgb = RGBColor(255, 255, 255)

    p_e2 = tf_e.add_paragraph()
    p_e2.text = "A fully responsive, production-ready campus placement web application"
    p_e2.font.size = Pt(16)
    p_e2.font.color.rgb = RGBColor(148, 163, 184)
    p_e2.space_before = Pt(6)

    # 4 Tech summary cards
    tech_cards = [
        ("Frontend Design", "Bootstrap 5 + Vanilla CSS\nMaterial Design Symbols\nMobile & Desktop Responsive\nChart.js Data Visualizations"),
        ("Backend Architecture", "PHP 8.2 (MVC Pattern)\nFront Controller Routing\nClean OOP Domain Models\nStrict Session Authentication"),
        ("Database Layer", "MySQL 8.0 Relational DB\nPDO Prepared Statements\nForeign Key Constraints\nIndexed Search Queries"),
        ("Security Features", "Bcrypt Password Hashing\nSQLi Immunity via PDO\nDuplicate Application Guards\nRole-Based Access Control")
    ]

    for idx, (th, tb) in enumerate(tech_cards):
        left = Inches(0.8 + idx * 2.95)
        t_card = s_end.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, Inches(2.6), Inches(2.8), Inches(3.2))
        t_card.fill.solid()
        t_card.fill.fore_color.rgb = RGBColor(30, 41, 59)
        t_card.line.color.rgb = RGBColor(51, 65, 85)

        tf = t_card.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.2)
        tf.margin_top = Inches(0.25)

        p = tf.paragraphs[0]
        p.text = th
        p.font.size = Pt(16)
        p.font.bold = True
        p.font.color.rgb = ACCENT

        p_desc = tf.add_paragraph()
        p_desc.text = tb
        p_desc.font.size = Pt(12)
        p_desc.font.color.rgb = RGBColor(203, 213, 225)
        p_desc.space_before = Pt(10)

    # Thank you / Q&A footer
    ft_end = s_end.shapes.add_textbox(Inches(0.8), Inches(6.2), Inches(11.7), Inches(0.6))
    p_f = ft_end.text_frame.paragraphs[0]
    p_f.text = "✨ Thank you! Open for Questions & Live Demonstration (localhost/crms/)"
    p_f.font.size = Pt(15)
    p_f.font.bold = True
    p_f.font.color.rgb = RGBColor(56, 189, 248)

    out_file = r"c:\Users\A D M I N\Desktop\project\mini project\CRMS_Website_Presentation.pptx"
    prs.save(out_file)
    print("Visual presentation saved at:", out_file)

if __name__ == "__main__":
    build_presentation()
