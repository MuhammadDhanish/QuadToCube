import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

def create_deck():
    prs = Presentation()
    # 16:9 Widescreen layout
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Brand Colors
    PRIMARY = RGBColor(10, 37, 64)       # Deep Navy
    SECONDARY = RGBColor(32, 96, 223)    # Bright Royal Blue
    ACCENT = RGBColor(14, 165, 233)      # Cyan / Sky
    BG_LIGHT = RGBColor(248, 250, 252)   # Slate 50
    TEXT_DARK = RGBColor(15, 23, 42)     # Slate 900
    TEXT_MUTED = RGBColor(100, 116, 139) # Slate 500
    CARD_BG = RGBColor(255, 255, 255)    # White
    CARD_BORDER = RGBColor(226, 232, 240)# Slate 200
    SUCCESS = RGBColor(16, 185, 129)     # Emerald

    def add_bg(slide, color=BG_LIGHT):
        shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
        shape.fill.solid()
        shape.fill.fore_color.rgb = color
        shape.line.color.rgb = color
        return shape

    def add_header(slide, title_text, category_text="CAMPUS RECRUITMENT MANAGEMENT SYSTEM"):
        # Category tag
        cat_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.5), Inches(11.7), Inches(0.4))
        tf_cat = cat_box.text_frame
        tf_cat.word_wrap = True
        p_cat = tf_cat.paragraphs[0]
        p_cat.text = category_text.upper()
        p_cat.font.size = Pt(11)
        p_cat.font.bold = True
        p_cat.font.color.rgb = SECONDARY
        p_cat.font.name = 'Arial'

        # Main Title
        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.85), Inches(11.7), Inches(0.7))
        tf = title_box.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = title_text
        p.font.size = Pt(26)
        p.font.bold = True
        p.font.color.rgb = PRIMARY
        p.font.name = 'Arial'

    # -------------------------------------------------------------
    # SLIDE 1: Title Slide (Dark Theme)
    # -------------------------------------------------------------
    s1 = prs.slides.add_slide(blank_layout)
    add_bg(s1, PRIMARY)

    # Accent decorative bar
    bar = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.8), Inches(1.2), Inches(0.08))
    bar.fill.solid()
    bar.fill.fore_color.rgb = ACCENT
    bar.line.fill.background()

    # Title & Subtitle box
    tbox = s1.shapes.add_textbox(Inches(0.8), Inches(2.2), Inches(11.7), Inches(3.2))
    tf = tbox.text_frame
    tf.word_wrap = True
    
    p = tf.paragraphs[0]
    p.text = "Campus Recruitment Management System"
    p.font.size = Pt(40)
    p.font.bold = True
    p.font.color.rgb = RGBColor(255, 255, 255)
    p.font.name = 'Arial'

    p2 = tf.add_paragraph()
    p2.text = "NexusCRMS: An Automated, Role-Based Placement & Drive Tracking Platform"
    p2.font.size = Pt(20)
    p2.font.color.rgb = RGBColor(186, 215, 255)
    p2.font.name = 'Arial'
    p2.space_before = Pt(12)

    p3 = tf.add_paragraph()
    p3.text = "Tech Stack: PHP (MVC) | MySQL | Bootstrap 5 | Chart.js | Apache (XAMPP)"
    p3.font.size = Pt(14)
    p3.font.color.rgb = RGBColor(148, 163, 184)
    p3.font.name = 'Arial'
    p3.space_before = Pt(24)

    # Footer note
    ft = s1.shapes.add_textbox(Inches(0.8), Inches(6.2), Inches(11.7), Inches(0.6))
    p_ft = ft.text_frame.paragraphs[0]
    p_ft.text = "Academic Mini Project Presentation | Department of Computer Science & Engineering"
    p_ft.font.size = Pt(12)
    p_ft.font.color.rgb = RGBColor(148, 163, 184)

    # -------------------------------------------------------------
    # SLIDE 2: Problem Statement & Motivation
    # -------------------------------------------------------------
    s2 = prs.slides.add_slide(blank_layout)
    add_bg(s2)
    add_header(s2, "The Problem: Why Traditional Placement Fails", "BACKGROUND & CHALLENGES")

    cards = [
        ("Manual & Error-Prone", "Universities rely heavily on scattered Excel sheets and Google Forms. Cross-verifying student CGPA, backlogs, and branch eligibility manually leads to critical human errors.", MSO_SHAPE.ROUNDED_RECTANGLE),
        ("Communication Gap", "Job circulars sent over WhatsApp groups and bulk emails get overlooked. Students miss stringent application deadlines and lack visibility on their application review stage.", MSO_SHAPE.ROUNDED_RECTANGLE),
        ("Administrative Overload", "Placement officers spend 80% of their time filtering resumes and compiling manual status logs rather than actively liaising with corporate recruiters and training students.", MSO_SHAPE.ROUNDED_RECTANGLE)
    ]

    for i, (head, body, shp) in enumerate(cards):
        left = Inches(0.8 + i * 4.0)
        card = s2.shapes.add_shape(shp, left, Inches(2.0), Inches(3.7), Inches(4.5))
        card.fill.solid()
        card.fill.fore_color.rgb = CARD_BG
        card.line.color.rgb = CARD_BORDER

        tf_c = card.text_frame
        tf_c.word_wrap = True
        tf_c.margin_left = Inches(0.3)
        tf_c.margin_right = Inches(0.3)
        tf_c.margin_top = Inches(0.4)
        
        p = tf_c.paragraphs[0]
        p.text = f"0{i+1}"
        p.font.size = Pt(28)
        p.font.bold = True
        p.font.color.rgb = SECONDARY

        p2 = tf_c.add_paragraph()
        p2.text = head
        p2.font.size = Pt(18)
        p2.font.bold = True
        p2.font.color.rgb = TEXT_DARK
        p2.space_before = Pt(14)

        p3 = tf_c.add_paragraph()
        p3.text = body
        p3.font.size = Pt(13)
        p3.font.color.rgb = TEXT_MUTED
        p3.space_before = Pt(12)

    # -------------------------------------------------------------
    # SLIDE 3: Proposed Solution & Core Value
    # -------------------------------------------------------------
    s3 = prs.slides.add_slide(blank_layout)
    add_bg(s3)
    add_header(s3, "NexusCRMS: The Centralized Placement Ecosystem", "OUR PROPOSED SOLUTION")

    features = [
        ("Smart Eligibility Engine", "Automatic real-time validation of Student CGPA, branch, and semester against job requirements. No ineligible submissions possible."),
        ("Transparent 4-Stage Tracker", "Live visual status stepper: Applied -> Under Review -> Shortlisted -> Selected with instant offer letter display."),
        ("Corporate & Drive Manager", "Full management of partner companies, job openings, CTC packages, criteria, and multi-round recruitment drives."),
        ("Executive Analytics & CSV Export", "Real-time placement KPI metrics, branch-wise placement donut charts, and one-click university report downloads.")
    ]

    for i, (title, desc) in enumerate(features):
        row = i // 2
        col = i % 2
        left = Inches(0.8 + col * 5.95)
        top = Inches(2.0 + row * 2.45)

        box = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, Inches(5.7), Inches(2.15))
        box.fill.solid()
        box.fill.fore_color.rgb = CARD_BG
        box.line.color.rgb = CARD_BORDER

        tf = box.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.3)
        tf.margin_top = Inches(0.25)

        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(16)
        p.font.bold = True
        p.font.color.rgb = PRIMARY

        p2 = tf.add_paragraph()
        p2.text = desc
        p2.font.size = Pt(13)
        p2.font.color.rgb = TEXT_MUTED
        p2.space_before = Pt(8)

    # -------------------------------------------------------------
    # SLIDE 4: System Architecture & Tech Stack
    # -------------------------------------------------------------
    s4 = prs.slides.add_slide(blank_layout)
    add_bg(s4)
    add_header(s4, "System Architecture: MVC Design Pattern", "TECHNICAL SPECIFICATIONS")

    tech_cols = [
        ("Presentation Layer", "Frontend UI", ["Responsive Bootstrap 5 & Vanilla CSS", "Material Symbols & Inter Typography", "Chart.js for interactive analytics", "Asynchronous fetch API & modal forms"]),
        ("Application Layer", "Backend PHP MVC", ["Front Controller (public/index.php)", "Dedicated Controllers (Auth, Job, App)", "Clean OOP Models (User, Job, Report)", "Session & CSRF security controls"]),
        ("Data Layer", "Relational Database", ["MySQL 8.0 InnoDB engine", "Referential integrity & foreign keys", "Prepared PDO statements (SQLi proof)", "Optimized indexing on search queries"])
    ]

    for i, (layer, sub, items) in enumerate(tech_cols):
        left = Inches(0.8 + i * 4.0)
        card = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, Inches(2.0), Inches(3.7), Inches(4.7))
        card.fill.solid()
        card.fill.fore_color.rgb = CARD_BG
        card.line.color.rgb = CARD_BORDER

        tf = card.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.25)
        tf.margin_top = Inches(0.3)

        p = tf.paragraphs[0]
        p.text = layer
        p.font.size = Pt(17)
        p.font.bold = True
        p.font.color.rgb = PRIMARY

        p2 = tf.add_paragraph()
        p2.text = sub
        p2.font.size = Pt(12)
        p2.font.color.rgb = SECONDARY
        p2.space_before = Pt(4)

        for item in items:
            pi = tf.add_paragraph()
            pi.text = f"•  {item}"
            pi.font.size = Pt(12)
            pi.font.color.rgb = TEXT_MUTED
            pi.space_before = Pt(10)

    # -------------------------------------------------------------
    # SLIDE 5: Student Workflow (End-to-End)
    # -------------------------------------------------------------
    s5 = prs.slides.add_slide(blank_layout)
    add_bg(s5)
    add_header(s5, "Student User Journey & Workflow", "USER EXPERIENCE")

    steps = [
        ("1. Profile Setup", "Student registers with branch, semester, verified CGPA, skills, and resume link."),
        ("2. Explore Drives", "Browse active openings; live badges show 'Eligible (Match 85%)' or specific reason if disqualified."),
        ("3. One-Click Apply", "Direct 1-click submission with cover notes. Guard checks prevent duplicate entries."),
        ("4. Real-Time Tracking", "Watch application journey live: Applied -> Under Review -> Shortlisted -> Placed!"),
        ("5. Offer Letter", "View confirmed offer letter reference code and coordinator feedback directly in dashboard.")
    ]

    for i, (st, desc) in enumerate(steps):
        top = Inches(2.0 + i * 0.95)
        box = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), top, Inches(11.7), Inches(0.8))
        box.fill.solid()
        box.fill.fore_color.rgb = CARD_BG
        box.line.color.rgb = CARD_BORDER

        tf = box.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.3)
        tf.margin_top = Inches(0.15)

        p = tf.paragraphs[0]
        p.text = st + ": "
        p.font.size = Pt(15)
        p.font.bold = True
        p.font.color.rgb = PRIMARY

        # Append desc in same paragraph
        run = p.add_run()
        run.text = desc
        run.font.bold = False
        run.font.color.rgb = TEXT_MUTED

    # -------------------------------------------------------------
    # SLIDE 6: Administrator & Coordinator Workflow
    # -------------------------------------------------------------
    s6 = prs.slides.add_slide(blank_layout)
    add_bg(s6)
    add_header(s6, "Placement Cell Administrative Controls", "ADMINISTRATIVE WORKFLOW")

    admin_cols = [
        ("1. Corporate Onboarding", "Create and maintain profiles for recruiting partners (Google, Microsoft, AWS, TCS, Infosys) including HR contacts and company tiers."),
        ("2. Placement Drive Launch", "Publish job drives specifying CTC package, job description, minimum CGPA, eligible branches, and deadline timestamp."),
        ("3. Candidate Review Board", "Filter applicant rosters by drive, review credentials, update workflow status across interview rounds, and assign official offer IDs."),
        ("4. Analytics & CSV Export", "Monitor live placement percentage, average salary metrics, and export sanitized accreditation-ready CSV reports for NAAC/NIRF.")
    ]

    for i, (title, desc) in enumerate(admin_cols):
        left = Inches(0.8 + i * 2.95)
        box = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, Inches(2.0), Inches(2.8), Inches(4.6))
        box.fill.solid()
        box.fill.fore_color.rgb = CARD_BG
        box.line.color.rgb = CARD_BORDER

        tf = box.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.2)
        tf.margin_top = Inches(0.3)

        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(16)
        p.font.bold = True
        p.font.color.rgb = SECONDARY

        p2 = tf.add_paragraph()
        p2.text = desc
        p2.font.size = Pt(13)
        p2.font.color.rgb = TEXT_MUTED
        p2.space_before = Pt(12)

    # -------------------------------------------------------------
    # SLIDE 7: Database Design & Key Entities
    # -------------------------------------------------------------
    s7 = prs.slides.add_slide(blank_layout)
    add_bg(s7)
    add_header(s7, "Relational Schema & Entity Design", "DATABASE ARCHITECTURE")

    tables = [
        ("students", "student_id (PK), name, email, password_hash, phone, course, semester, cgpa, skills, resume_link"),
        ("admins", "admin_id (PK), name, email, password_hash, role ('superadmin', 'coordinator')"),
        ("companies", "company_id (PK), company_name, contact_email, website, industry, location, description"),
        ("jobs", "job_id (PK), company_id (FK), title, salary_package, min_cgpa, allowed_courses, deadline, status"),
        ("applications", "application_id (PK), job_id (FK), student_id (FK), applied_date, status, notes"),
        ("selections", "selection_id (PK), application_id (FK), result, offer_letter_ref, remarks")
    ]

    for i, (tname, fields) in enumerate(tables):
        row = i // 2
        col = i % 2
        left = Inches(0.8 + col * 5.95)
        top = Inches(2.0 + row * 1.6)

        box = s7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, Inches(5.7), Inches(1.4))
        box.fill.solid()
        box.fill.fore_color.rgb = CARD_BG
        box.line.color.rgb = CARD_BORDER

        tf = box.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.25)
        tf.margin_top = Inches(0.18)

        p = tf.paragraphs[0]
        p.text = f"TABLE: {tname}"
        p.font.size = Pt(15)
        p.font.bold = True
        p.font.color.rgb = PRIMARY

        p2 = tf.add_paragraph()
        p2.text = fields
        p2.font.size = Pt(11)
        p2.font.color.rgb = TEXT_MUTED
        p2.space_before = Pt(4)

    # -------------------------------------------------------------
    # SLIDE 8: Key Features & Technical Highlights
    # -------------------------------------------------------------
    s8 = prs.slides.add_slide(blank_layout)
    add_bg(s8)
    add_header(s8, "Key Highlights & Security Implementations", "ENGINEERING MERITS")

    highlights = [
        ("Security: Bcrypt & Prepared Statements", "Zero plain-text passwords stored. All database queries execute via PDO prepared parameters to completely eradicate SQL Injection vulnerabilities."),
        ("Automated Validation Engine", "System automatically verifies student CGPA >= job min_cgpa, checks branch in allowed_courses list, and ensures current timestamp < deadline."),
        ("Duplicate Submission Guard", "Database unique constraints and controller guard clauses prevent accidental double submissions for the same opening."),
        ("Fast XAMPP & Apache Integration", "Seamless zero-configuration deployment on localhost Apache with URL rewrites and root forwarder.")
    ]

    for i, (t, d) in enumerate(highlights):
        top = Inches(2.0 + i * 1.2)
        box = s8.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), top, Inches(11.7), Inches(1.05))
        box.fill.solid()
        box.fill.fore_color.rgb = CARD_BG
        box.line.color.rgb = CARD_BORDER

        tf = box.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.3)
        tf.margin_top = Inches(0.15)

        p = tf.paragraphs[0]
        p.text = t
        p.font.size = Pt(16)
        p.font.bold = True
        p.font.color.rgb = PRIMARY

        p2 = tf.add_paragraph()
        p2.text = d
        p2.font.size = Pt(12)
        p2.font.color.rgb = TEXT_MUTED
        p2.space_before = Pt(4)

    # -------------------------------------------------------------
    # SLIDE 9: Live Demo Walkthrough
    # -------------------------------------------------------------
    s9 = prs.slides.add_slide(blank_layout)
    add_bg(s9)
    add_header(s9, "Live Demonstration Walkthrough", "DEMO PLAN")

    steps_demo = [
        ("Step 1: Admin Dashboard", "Log in as admin@crms.edu | Review real-time KPIs, branch placement donut chart, and drive metrics."),
        ("Step 2: Post New Drive", "Navigate to Jobs -> Create new drive with criteria (e.g. 7.5 CGPA, CSE/IT, 12 LPA)."),
        ("Step 3: Student Experience", "Log in as rahul.sharma@crms.edu | Observe smart eligibility badges and apply with 1 click."),
        ("Step 4: Shortlisting & Offer", "Switch back to coordinator | Shortlist candidate, advance stage, and grant offer letter reference.")
    ]

    for i, (st, desc) in enumerate(steps_demo):
        top = Inches(2.0 + i * 1.2)
        box = s9.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), top, Inches(11.7), Inches(1.05))
        box.fill.solid()
        box.fill.fore_color.rgb = CARD_BG
        box.line.color.rgb = CARD_BORDER

        tf = box.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.3)
        tf.margin_top = Inches(0.15)

        p = tf.paragraphs[0]
        p.text = st
        p.font.size = Pt(16)
        p.font.bold = True
        p.font.color.rgb = SECONDARY

        p2 = tf.add_paragraph()
        p2.text = desc
        p2.font.size = Pt(13)
        p2.font.color.rgb = TEXT_MUTED
        p2.space_before = Pt(4)

    # -------------------------------------------------------------
    # SLIDE 10: Conclusion & Future Scope (Dark Theme)
    # -------------------------------------------------------------
    s10 = prs.slides.add_slide(blank_layout)
    add_bg(s10, PRIMARY)

    bar2 = s10.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.5), Inches(1.2), Inches(0.08))
    bar2.fill.solid()
    bar2.fill.fore_color.rgb = ACCENT
    bar2.line.fill.background()

    tbox = s10.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(11.7), Inches(1.0))
    tf = tbox.text_frame
    p = tf.paragraphs[0]
    p.text = "Conclusion & Future Enhancements"
    p.font.size = Pt(32)
    p.font.bold = True
    p.font.color.rgb = RGBColor(255, 255, 255)

    futures = [
        ("Resume Parser (AI/NLP)", "Automate skill extraction from PDF resumes to match candidate suitability score."),
        ("Automated Email / SMS Alerts", "Notify students instantaneously via SendGrid or Twilio when shortlisted or scheduled for interviews."),
        ("Online Assessment Module", "Integrated coding / aptitude testing platform directly within the placement cell portal.")
    ]

    for i, (title, desc) in enumerate(futures):
        left = Inches(0.8 + i * 4.0)
        card = s10.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, Inches(3.0), Inches(3.7), Inches(3.2))
        card.fill.solid()
        card.fill.fore_color.rgb = RGBColor(15, 33, 60)
        card.line.color.rgb = RGBColor(30, 58, 102)

        tf = card.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.25)
        tf.margin_top = Inches(0.3)

        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(16)
        p.font.bold = True
        p.font.color.rgb = ACCENT

        p2 = tf.add_paragraph()
        p2.text = desc
        p2.font.size = Pt(13)
        p2.font.color.rgb = RGBColor(203, 213, 225)
        p2.space_before = Pt(10)

    # Save presentation
    output_path = r"c:\Users\A D M I N\Desktop\project\mini project\CRMS_Project_Presentation.pptx"
    prs.save(output_path)
    print("Presentation saved successfully at:", output_path)

if __name__ == "__main__":
    create_deck()
